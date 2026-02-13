"""
Config Loader Module
Reads and parses the master workbook into structured JSON objects.
Excel is READ-ONLY - we only parse values, never execute formulas.
Uses table names and sheet names, not fixed cell positions.
"""

import openpyxl
from typing import Dict, List, Any, Optional
from pathlib import Path
import logging

logger = logging.getLogger(__name__)


class ConfigLoader:
    """Loads configuration from master Excel workbook using table names"""
    
    def __init__(self, workbook_path: str):
        self.workbook_path = Path(workbook_path)
        self.workbook = None
        self.config = {
            "schema": [],
            "pricing_flow": [],
            "rules": [],
            "rate_library_map": [],
            "worked_examples": []
        }
        
    def load(self) -> Dict[str, Any]:
        """Load all configuration from workbook using table names"""
        if not self.workbook_path.exists():
            raise FileNotFoundError(f"Master workbook not found: {self.workbook_path}")
        
        logger.info(f"Loading master workbook: {self.workbook_path}")
        # data_only=True ensures we read VALUES not formulas (read-only principle)
        self.workbook = openpyxl.load_workbook(self.workbook_path, data_only=True)
        
        # Load each table by name
        self.config["schema"] = self._load_table("tbl_SCHEMA", "__SCHEMA__")
        self.config["pricing_flow"] = self._load_table("tbl_Pricing_Flow", "Pricing_Flow")
        self.config["rules"] = self._load_table("tbl_Rules", "Rules")
        self.config["rate_library_map"] = self._load_table("tbl_Rate_Library_Map", "Rate_Library_Map")
        
        logger.info("Master workbook loaded successfully")
        return self.config
    
    def _load_table(self, table_name: str, sheet_name: str) -> List[Dict[str, Any]]:
        """
        Load data from a named Excel table using worksheet.tables API.
        NO fixed cell positions, NO header scanning - uses table ref range.
        """
        # Find the table across all worksheets
        for sheet in self.workbook.worksheets:
            if hasattr(sheet, 'tables') and table_name in sheet.tables:
                table = sheet.tables[table_name]
                return self._parse_table_data(sheet, table, table_name)
        
        # Table not found - this is an error
        raise ValueError(
            f"Required table '{table_name}' not found in master workbook. "
            f"Expected tables: tbl_SCHEMA, tbl_Pricing_Flow, tbl_Rules, tbl_Rate_Library_Map"
        )
    
    def _parse_table_data(self, sheet, table, table_name: str) -> List[Dict[str, Any]]:
        """
        Parse Excel table data using table.ref range.
        First row in range is headers, subsequent rows are data.
        """
        # Get table range (e.g., "A1:F100")
        data_range = table.ref
        logger.info(f"Reading table '{table_name}' from range {data_range}")
        
        # Parse the table data
        data = []
        headers = None
        
        # Iterate over cells in the table range
        for row_idx, row in enumerate(sheet[data_range], start=1):
            row_values = [cell.value for cell in row]
            
            # First row is headers
            if headers is None:
                headers = []
                for i, cell_value in enumerate(row_values):
                    if cell_value:
                        headers.append(str(cell_value).strip())
                    else:
                        headers.append(f"col_{i}")
                logger.debug(f"Table '{table_name}' headers: {headers}")
                continue
            
            # Skip empty rows
            if not any(row_values):
                continue
            
            # Parse data row
            row_data = {}
            for i, cell_value in enumerate(row_values):
                if i < len(headers):
                    # Handle None values from uncached formulas
                    if cell_value is None:
                        logger.warning(
                            f"Table '{table_name}' row {row_idx}: cell in column '{headers[i]}' is None. "
                            f"This may indicate an uncached formula. Excel must be saved with values cached."
                        )
                    row_data[headers[i]] = cell_value
            
            if any(row_data.values()):  # Only add if row has some data
                data.append(row_data)
        
        logger.info(f"Loaded {len(data)} rows from table '{table_name}'")
        return data
    
    def get_schema_for_entity(self, entity: str) -> List[Dict[str, Any]]:
        """Get schema fields for a specific entity"""
        return [field for field in self.config["schema"] if field.get("entity") == entity]
    
    def get_pricing_flow_for_stage(self, stage: str) -> Optional[Dict[str, Any]]:
        """Get pricing flow rule for a specific stage"""
        for flow in self.config["pricing_flow"]:
            if flow.get("project_stage") == stage:
                return flow
        return None
    
    def get_rules_for_condition(self, **conditions) -> List[Dict[str, Any]]:
        """Get rules matching specific conditions"""
        matching_rules = []
        for rule in self.config["rules"]:
            matches = True
            for key, value in conditions.items():
                if rule.get(key) != value:
                    matches = False
                    break
            if matches:
                matching_rules.append(rule)
        return matching_rules


# Global config instance
_config_loader: Optional[ConfigLoader] = None


def get_config_loader() -> Optional[ConfigLoader]:
    """Get global config loader instance"""
    return _config_loader


def initialize_config(workbook_path: str) -> ConfigLoader:
    """Initialize global config loader"""
    global _config_loader
    _config_loader = ConfigLoader(workbook_path)
    _config_loader.load()
    return _config_loader
