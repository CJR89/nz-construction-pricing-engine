import pytest
from app.core.exceptions import ConfigError
from app.services.config_loader import ConfigLoader, REQUIRED_TABLES
from unittest.mock import Mock, MagicMock, patch
import openpyxl


def test_missing_table_raises_config_error():
    """Test that missing required tables raise ConfigError with detailed message"""
    # Create a mock workbook with only some tables
    mock_workbook = Mock()
    mock_sheet = Mock()
    mock_sheet.title = "TestSheet"
    mock_sheet.tables = {"tbl_SCHEMA": Mock()}  # Only tbl_SCHEMA exists
    
    mock_workbook.worksheets = [mock_sheet]
    
    loader = ConfigLoader("/fake/path/test.xlsm")
    loader.workbook_path = Mock()
    loader.workbook_path.exists.return_value = True
    loader.workbook_path.name = "test.xlsm"
    loader.workbook = mock_workbook
    
    # Should raise ConfigError listing missing tables
    with pytest.raises(ConfigError) as exc_info:
        loader._validate_required_tables()
    
    error_message = str(exc_info.value)
    # Check that error message contains missing tables
    assert "tbl_Pricing_Flow" in error_message
    assert "tbl_Rules" in error_message
    assert "tbl_Rate_Library_Map" in error_message
    assert "test.xlsm" in error_message
    assert "TestSheet" in error_message


def test_all_tables_present_no_error():
    """Test that no error is raised when all required tables exist"""
    mock_workbook = Mock()
    mock_sheet = Mock()
    mock_sheet.title = "ConfigSheet"
    mock_sheet.tables = {
        "tbl_SCHEMA": Mock(),
        "tbl_Pricing_Flow": Mock(),
        "tbl_Rules": Mock(),
        "tbl_Rate_Library_Map": Mock()
    }
    
    mock_workbook.worksheets = [mock_sheet]
    
    loader = ConfigLoader("/fake/path/test.xlsm")
    loader.workbook = mock_workbook
    
    # Should not raise any error
    loader._validate_required_tables()  # No exception means success


def test_config_error_on_missing_tbl_pricing_flow():
    """Specific test for missing tbl_Pricing_Flow"""
    mock_workbook = Mock()
    mock_sheet = Mock()
    mock_sheet.title = "Sheet1"
    mock_sheet.tables = {
        "tbl_SCHEMA": Mock(),
        "tbl_Rules": Mock(),
        "tbl_Rate_Library_Map": Mock()
        # tbl_Pricing_Flow is missing
    }
    
    mock_workbook.worksheets = [mock_sheet]
    
    loader = ConfigLoader("/fake/path/test.xlsm")
    loader.workbook_path = Mock()
    loader.workbook_path.exists.return_value = True
    loader.workbook_path.name = "test.xlsm"
    loader.workbook = mock_workbook
    
    with pytest.raises(ConfigError) as exc_info:
        loader._validate_required_tables()
    
    assert "tbl_Pricing_Flow" in str(exc_info.value)


def test_config_error_on_missing_tbl_rules():
    """Specific test for missing tbl_Rules"""
    mock_workbook = Mock()
    mock_sheet = Mock()
    mock_sheet.title = "Sheet1"
    mock_sheet.tables = {
        "tbl_SCHEMA": Mock(),
        "tbl_Pricing_Flow": Mock(),
        "tbl_Rate_Library_Map": Mock()
        # tbl_Rules is missing
    }
    
    mock_workbook.worksheets = [mock_sheet]
    
    loader = ConfigLoader("/fake/path/test.xlsm")
    loader.workbook_path = Mock()
    loader.workbook_path.exists.return_value = True
    loader.workbook_path.name = "test.xlsm"
    loader.workbook = mock_workbook
    
    with pytest.raises(ConfigError) as exc_info:
        loader._validate_required_tables()
    
    assert "tbl_Rules" in str(exc_info.value)


def test_config_error_lists_searched_sheets():
    """Test that error message lists which sheets were searched"""
    mock_workbook = Mock()
    
    sheet1 = Mock()
    sheet1.title = "ConfigSheet"
    sheet1.tables = {}
    
    sheet2 = Mock()
    sheet2.title = "DataSheet"
    sheet2.tables = {}
    
    mock_workbook.worksheets = [sheet1, sheet2]
    
    loader = ConfigLoader("/fake/path/test.xlsm")
    loader.workbook_path = Mock()
    loader.workbook_path.exists.return_value = True
    loader.workbook_path.name = "test.xlsm"
    loader.workbook = mock_workbook
    
    with pytest.raises(ConfigError) as exc_info:
        loader._validate_required_tables()
    
    error_message = str(exc_info.value)
    assert "ConfigSheet" in error_message
    assert "DataSheet" in error_message
    assert "Searched sheets:" in error_message
